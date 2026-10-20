"""Bacaan dubbing: naskah narasi per adegan dengan huruf besar, nyaman dibaca dari HP atau cetakan A5.
Pakai: python buat-bacaan.py keluaran.docx"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from naskah import S, sec
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

NAVY, EMAS, ABU = RGBColor(0x1E, 0x3A, 0x6E), RGBColor(0xA8, 0x86, 0x2E), RGBColor(0x70, 0x70, 0x70)
d = Document(); sc = d.sections[0]
sc.page_width, sc.page_height = Cm(14.8), Cm(21)
sc.left_margin = sc.right_margin = Cm(1.4); sc.top_margin = sc.bottom_margin = Cm(1.4)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(12); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def para(t, b=False, sz=12, warna=None, i=False, after=4, al=None, ls=None):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; r.italic = i; r.font.size = Pt(sz)
    if warna: r.font.color.rgb = warna
    p.paragraph_format.space_after = Pt(after)
    if ls: p.paragraph_format.line_spacing = ls
    if al is not None: p.alignment = al
    return p

para("BACAAN DUBBING", True, 18, NAVY, al=WD_ALIGN_PARAGRAPH.CENTER, after=2)
para("Video Prokopim Hibot · Kaltara Innovation Awards 2026 · 4:55", False, 11, al=WD_ALIGN_PARAGRAPH.CENTER, after=10)
for t in ["Baca teks hitam besar persis seperti tertulis; angka sudah dieja.",
          "Satu rekaman per adegan, dengan jeda diam 1 detik di awal dan akhir.",
          "Waktu di judul adalah posisi di video; jangan melebihi batas detiknya.",
          "Tulisan abu-abu miring adalah arahan, tidak dibaca."]:
    p = d.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(2)
for no, a, b, sub, teks, cat in S:
    d.add_page_break()
    para(f"ADEGAN {no}", True, 20, NAVY, after=0)
    para(f"{a}–{b}  ·  maks {sec(b) - sec(a)} detik  ·  {sub}", False, 11, EMAS, after=12)
    if no == "13":
        para("TIDAK ADA NARASI", True, 20, RGBColor(0xB9, 0x1C, 0x1C), after=6)
        para("Suara asli Wali Kota diputar. Dubber diam.", False, 16, after=12, ls=1.3)
    else:
        for k in [x for x in teks.replace(". ", ".\n").split("\n") if x.strip()]:
            para(k.strip(), False, 19, after=10, ls=1.3)
    para(cat, False, 11.5, ABU, i=True, after=0)
d.save(sys.argv[1])
