import sys; sys.path.insert(0, ".")
from naskah import S, sec
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
sc = d.sections[0]
sc.page_width, sc.page_height = Cm(21), Cm(29.7)
sc.left_margin = sc.right_margin = Cm(2); sc.top_margin = sc.bottom_margin = Cm(2)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def para(t="", b=False, sz=None, al=None, after=4, i=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; r.italic = i
    if sz: r.font.size = Pt(sz)
    if al: p.alignment = al
    p.paragraph_format.space_after = Pt(after)
    return p
def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

para("LEMBAR NASKAH PENGISI SUARA", True, 16, WD_ALIGN_PARAGRAPH.CENTER, 2)
para("Video Inovasi Prokopim Hibot — Kaltara Innovation Awards 2026", False, 11, WD_ALIGN_PARAGRAPH.CENTER, 2)
para("Bagian Protokol dan Komunikasi Pimpinan, Setda Kota Tarakan", False, 10, WD_ALIGN_PARAGRAPH.CENTER, 12)

para("Petunjuk rekaman", True, 12, after=4)
for t in [
 "Rekam di ruang tenang, tanpa kipas atau AC yang berisik. Jarak mulut ke mikrofon sekitar 20 cm; bila memakai HP, letakkan menghadap mulut.",
 "Satu berkas per adegan, diberi nama sesuai nomor: adegan-01, adegan-02, dan seterusnya. Format WAV atau M4A.",
 "Rekam dua kali per adegan; kirim keduanya. Beri jeda diam 1 detik di awal dan di akhir setiap berkas.",
 "Tempo santai dan jelas, sekitar 110–130 kata per menit. Durasi di tabel adalah batas maksimum; lebih pendek tidak apa-apa, lebih panjang tidak boleh.",
 "Angka sudah ditulis dengan huruf agar dibaca persis. Baca apa adanya, jangan diubah atau dibulatkan.",
 "\"Prokopim Hibot\" dibaca sebagai kata, bukan dieja. \"RK\" ditulis \"Rencana Kegiatan\"; \"SOP\" dibaca \"es-o-pe\".",
 "Nada: resmi namun hangat, seperti menjelaskan kepada rekan kerja. Hindari gaya iklan.",
]:
    p = d.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(2)
para("", after=6)

t = d.add_table(rows=1, cols=4); t.style = "Table Grid"; t.autofit = False
W = [Cm(1.2), Cm(2.6), Cm(8.7), Cm(4.5)]
for c, h in zip(t.rows[0].cells, ["No.", "Waktu / batas", "Naskah (dibaca persis)", "Catatan membaca"]):
    c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1E3A6E")
for n, a, b, sub, txt, cat in S:
    dur = sec(b) - sec(a); kata = len(txt.split())
    row = t.add_row().cells
    row[0].text = ""; r = row[0].paragraphs[0].add_run(n); r.bold = True
    row[1].text = ""; p = row[1].paragraphs[0]; p.add_run(f"{a}–{b}\n").font.size = Pt(10); r = p.add_run(f"maks {dur} dtk\n{sub}"); r.font.size = Pt(9); r.font.color.rgb = RGBColor(90, 90, 90)
    row[2].text = ""; r = row[2].paragraphs[0].add_run(txt); r.font.size = Pt(12)
    row[3].text = ""; r = row[3].paragraphs[0].add_run(cat); r.font.size = Pt(10); r.italic = True
for row in t.rows:
    for c, w in zip(row.cells, W): c.width = w
    trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
for col, w in zip(t.columns, W): col.width = w
grid = t._tbl.tblGrid
for gc, w in zip(grid.findall(qn("w:gridCol")), W): gc.set(qn("w:w"), str(int(w.twips)))
hdr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); hdr.append(th)

para("", after=6)
para("Yang dikirim kembali: seluruh berkas adegan, ke Admin Rencana Kegiatan atau Kabag Prokopim. Bila ada kata yang sulit diucapkan atau kalimat yang terasa terlalu panjang, tandai nomor adegannya dan sampaikan; naskah bisa dipersingkat.", i=True, after=2)
d.save(sys.argv[1])
