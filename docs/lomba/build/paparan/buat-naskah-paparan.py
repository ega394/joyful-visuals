"""Lembar naskah paparan 10 menit + panduan latihan (DOCX). Pakai: python buat-naskah-paparan.py keluaran.docx"""
import os, re, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from naskah_paparan import SLIDE, TANYA, KIAT, kata
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY, EMAS, ABU = RGBColor(0x1E, 0x3A, 0x6E), RGBColor(0xA8, 0x86, 0x2E), RGBColor(0x80, 0x80, 0x80)
d = Document(); sc = d.sections[0]
sc.page_width, sc.page_height = Cm(21), Cm(29.7)
sc.left_margin = sc.right_margin = Cm(2); sc.top_margin = sc.bottom_margin = Cm(1.8)
st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

def para(t="", b=False, sz=None, al=None, after=4, i=False, warna=None):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; r.italic = i
    if sz: r.font.size = Pt(sz)
    if warna: r.font.color.rgb = warna
    if al is not None: p.alignment = al
    p.paragraph_format.space_after = Pt(after)
    return p
def shade(c, h):
    p = c._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), h); p.append(s)
def tabel(baris, lebar, kepala):
    t = d.add_table(rows=1, cols=len(lebar)); t.style = "Table Grid"; t.autofit = False
    for c, h in zip(t.rows[0].cells, kepala):
        c.text = ""; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1E3A6E")
    for b in baris:
        cs = t.add_row().cells
        for c, v in zip(cs, b):
            c.text = ""; r = c.paragraphs[0].add_run(v); r.font.size = Pt(10.5)
    for row in t.rows:
        for c, w in zip(row.cells, lebar): c.width = w
        x = OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(x)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), lebar): gc.set(qn("w:w"), str(int(w.twips)))
    return t
def ucapan(teks, ukuran=14):
    """Kalimat biasa hitam; bagian [[...]] abu-abu miring (boleh dilewati)."""
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(7); p.paragraph_format.line_spacing = 1.25
    for bag in re.split(r"(\[\[.*?\]\])", teks):
        if not bag: continue
        opsional = bag.startswith("[[")
        r = p.add_run(bag[2:-2] if opsional else bag); r.font.size = Pt(ukuran)
        if opsional: r.italic = True; r.font.color.rgb = ABU

para("NASKAH PAPARAN TAHAP II", True, 17, WD_ALIGN_PARAGRAPH.CENTER, 2, warna=NAVY)
para("Prokopim Hibot · Kaltara Innovation Awards 2026", False, 12, WD_ALIGN_PARAGRAPH.CENTER, 2)
para("Kategori Inovasi Terapan – ASN Pemerintah Kabupaten/Kota", True, 11, WD_ALIGN_PARAGRAPH.CENTER, 2, warna=EMAS)
para("Penilaian 22 Oktober 2026 · Ruang B, Hotel Luminor Tanjung Selor · Penyaji: Anugrah Yega Pranatha, M.Si.", False, 10, WD_ALIGN_PARAGRAPH.CENTER, 12)

para("Cara memakai lembar ini", True, 13, after=3, warna=NAVY)
for t in ["Waktu paparan dari panitia 10 menit, termasuk demonstrasi. Naskah ini dirancang selesai sekitar 9:15, menyisakan cadangan 45 detik.",
          "Teks hitam diucapkan. Teks abu-abu miring boleh dilewati bila waktu mepet. Kotak 'Isyarat' adalah gerak dan nada, tidak diucapkan.",
          "Lembar ini untuk latihan. Di depan juri, cukup pegang halaman 'Pos waktu' dan kalimat pertama tiap slide.",
          "Naskah yang sama tersimpan sebagai catatan pembicara di PowerPoint (tampilan Presenter)."]:
    p = d.add_paragraph(t, style="List Bullet"); p.paragraph_format.space_after = Pt(2)

para("", after=4)
para("Pos waktu", True, 13, after=3, warna=NAVY)
tabel([[f"{s['mulai']}–{s['selesai']}", str(s["no"]), s["judul"], s["ucap"][0].split(" — ")[0][:70] + ("…" if len(s["ucap"][0]) > 70 else "")] for s in SLIDE[:13]],
      [Cm(2.4), Cm(1.4), Cm(4.4), Cm(8.8)], ["Waktu", "Slide", "Bagian", "Kalimat pembuka"])
para("Pos wajib: slide 2 paling lambat 0:30 · demo mulai 2:40 · slide 7 mulai 5:40 · penutup mulai 8:50.", True, 10.5, after=0, warna=RGBColor(0xB9, 0x1C, 0x1C))

for s in SLIDE:
    d.add_page_break() if s["no"] in (1, 6, 7, 14) else None
    k = sum(kata(u) for u in s["ucap"])
    h = para(f"SLIDE {s['no']} · {s['judul'].upper()}", True, 13, after=1, warna=NAVY)
    h.paragraph_format.space_before = Pt(10); h.paragraph_format.keep_with_next = True
    info = f"{s['mulai']}–{s['selesai']}" + (f" · ±{k} kata" if s["no"] not in (6, 14) else "")
    para(info, False, 10, after=4, warna=EMAS).paragraph_format.keep_with_next = True
    for u in s["ucap"]:
        ucapan(u, 13 if s["no"] == 6 else 14)
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"; c = t.rows[0].cells[0]; shade(c, "F3EEDD")
    c.text = ""; r = c.paragraphs[0].add_run("Isyarat: "); r.bold = True; r.font.size = Pt(10.5)
    r = c.paragraphs[0].add_run(s["isyarat"]); r.font.size = Pt(10.5); r.italic = True

d.add_page_break()
para("TANYA JAWAB: PERTANYAAN YANG MUNGKIN MUNCUL", True, 14, after=4, warna=NAVY)
para("Pola menjawab: terima kasih → ulangi inti pertanyaan → jawaban → bukti → manfaat. Paling lama 30 detik per jawaban.", False, 10.5, after=6, i=True)
tabel([[q, a] for q, a in TANYA], [Cm(5.2), Cm(11.8)], ["Pertanyaan", "Inti jawaban"])

d.add_page_break()
para("PANDUAN LATIHAN DAN PENYAMPAIAN", True, 14, after=6, warna=NAVY)
for j, isi in KIAT:
    p = d.add_paragraph(); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(j + ". "); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = NAVY
    r = p.add_run(isi); r.font.size = Pt(11.5)
para("", after=4)
para("Lembar latihan stopwatch", True, 12, after=3, warna=NAVY)
tabel([["Latihan " + str(i), "", "", "", ""] for i in range(1, 6)], [Cm(3), Cm(3.5), Cm(3.5), Cm(3.5), Cm(3.5)],
      ["", "Slide 2 (≤0:30)", "Demo (2:40)", "Slide 7 (5:40)", "Selesai (≤9:30)"])
d.save(sys.argv[1])
